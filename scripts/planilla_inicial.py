# Arma planilla/Planilla_inicial_auditoria.xlsx: las hojas de datos de la app, ya cargadas con
# los números del Excel de Senderos II (06 a 10/2026) y La Huella (08 a 10/2026).
# Uso: python3 scripts/planilla_inicial.py
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

esq_ts = open('src/data/esquema.ts', encoding='utf8').read()
copia = json.load(open('src/data/copia.json', encoding='utf8'))
hojas = re.findall(r"^  ([A-Z_]+): \{\n    id: '([^']+)',\n    publica: (?:true|false),\n    descripcion: '((?:[^'\\]|\\.)*)',\n    columnas: \[(.*?)\],", esq_ts, re.S | re.M)
wb = Workbook()
le = wb.active; le.title = 'LEEME'
txt = [
    'Planilla de datos de la Auditoría de expensas',
    '',
    'Cada hoja es una tabla: una fila por registro, sin fórmulas. La app lee estas hojas y arma todo el análisis.',
    'No cambies los nombres de las hojas ni de las columnas (fila 1). Sí podés agregar o corregir filas.',
    'Importes en pesos con decimales (1234.56). Meses como AAAA-MM. Fechas como AAAA-MM-DD. SI/NO en mayúsculas.',
    '',
    'Contiene: Senderos II de 06 a 10/2026 y La Huella de 08 a 10/2026, con la liquidación correcta, el respaldo de cada gasto,',
    'la matriz de riesgo, la cronología, las notas y los ahorros posibles del Excel por rubro.',
    'No contiene nombres de propietarios, deuda por lote ni CUIT de personas físicas.',
]
for i, t in enumerate(txt, 1):
    le.cell(i, 1, t)
le['A1'].font = Font(bold=True, size=14); le.column_dimensions['A'].width = 120
n = 0
for nombre, _id, desc, cols in hojas:
    columnas = re.findall(r"'([^']+)'", cols)
    ws = wb.create_sheet(nombre)
    for j, c in enumerate(columnas, 1):
        x = ws.cell(1, j, c); x.font = Font(bold=True); x.fill = PatternFill('solid', fgColor='DCE4F0')
        ws.column_dimensions[x.column_letter].width = max(12, min(40, len(c) + 4))
    ws.freeze_panes = 'A2'
    for i, f in enumerate(copia['sheets'].get(nombre, []), 2):
        for j, c in enumerate(columnas, 1):
            v = f.get(c, '')
            if v is None: v = ''
            cell = ws.cell(i, j, v)
            if isinstance(v, str): cell.number_format = '@'
        n += 1
    ws.cell(1, 1).comment = None
wb.save('planilla/Planilla_inicial_auditoria.xlsx')
print(len(hojas), 'hojas,', n, 'filas')
